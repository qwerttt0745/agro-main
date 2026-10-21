"""Versioned official register snapshots, exact provenance, no application advice."""
import hashlib
import io
import json
import re
import sqlite3
import threading
import unicodedata
from datetime import datetime, timezone, date
from zoneinfo import ZoneInfo
from urllib.parse import urljoin, urlparse
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from app import config

DB=config.DATA/'registers.sqlite3'
LOCK=threading.Lock()
SOURCES={
 'pesticides':{'title':'Реєстр пестицидів і агрохімікатів','index':'https://me.gov.ua/Documents/List?id=00aa19f1-7ab1-4305-b69c-08bd6064acf4&lang=uk-UA&tag=DerzhavniiRestrPestitsidivIAgrokhimikativDozvolenikhDoVikoristanniaVUkraini','match':'державний реєстр пестицидів'},
 'varieties':{'title':'Реєстр сортів рослин','index':'https://me.gov.ua/Documents/List?id=b5620a37-4f55-4689-b2a3-a0f8c07f739f&lang=uk-UA&tag=RestriUSferiOkhoroniPravNaSortiRoslin','match':'державний реєстр сортів, придатних'}
}


def normalized(value):
    return ' '.join(unicodedata.normalize('NFKC',str(value or '')).casefold().split())


def connect():
    c=sqlite3.connect(DB,timeout=30);c.row_factory=sqlite3.Row
    c.execute('CREATE TABLE IF NOT EXISTS snapshots(kind TEXT PRIMARY KEY, metadata TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS records(kind TEXT,sheet TEXT,row_no INTEGER,name TEXT,search TEXT,payload TEXT)')
    c.execute('CREATE INDEX IF NOT EXISTS records_kind ON records(kind)')
    return c


def cell(value):
    if isinstance(value,datetime):return value.date().isoformat()
    return str(value).strip() if value is not None else ''


def parse_book(content,kind):
    book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    records=[]
    try:
        for sheet in book:
            rows=sheet.iter_rows(values_only=True)
            headers=[normalized(v) for v in next(rows)]
            if kind=='varieties':
                if 'назва сорту (українською)' not in headers:continue
                columns={'name':'назва сорту (українською)','crop':'ботанічний таксон (українською)','number':'номер заявки','zone':'рекомендована зона для вирощування','maturity':'група стиглості (скорочено)','year':'рік державної реєстрації','notes':'статус сорту'}
            elif 'registration_number' in headers:
                columns={'name':'name','number':'registration_number','certificate':'registration_certificate','scope':'scope','expires':'expiration_date','notes':'примітки'}
            else:
                columns={'name':'торгова назва препарату та позначення його препаративної форми','number':'реєстраційний номер препарату','certificate':'серія','ingredient':'найменування та вміст діючої речовини','expires':'строк реєстрації препарату','cancelled':'дата скасування (призупинення) державної реєстрації препарату','notes':'примітки'}
                scope=next((h for h in headers if h.startswith('сфера застосування')),None)
                if scope:columns['scope']=scope
            if any(v not in headers for v in columns.values()):raise ValueError('Register column schema changed')
            indexes={k:headers.index(v) for k,v in columns.items()}
            for row_no,row in enumerate(rows,2):
                payload={k:cell(row[i]) for k,i in indexes.items()}
                if not payload['name']:continue
                if any(v.startswith('=') for v in payload.values()):raise ValueError('Formula requires manual review')
                if kind=='varieties':
                    payload['section']='excluded' if 'виключ' in sheet.title.lower() else 'listed'
                    if 'дата припинення чинності державної реєстрації' in headers:
                        payload['excluded_on']=cell(row[headers.index('дата припинення чинності державної реєстрації')])
                records.append((kind,sheet.title,row_no,payload['name'],normalized(' '.join([payload['name'],payload.get('number',''),payload.get('ingredient',''),payload.get('crop','')])) ,json.dumps(payload,ensure_ascii=False)))
    finally:book.close()
    if not records:raise ValueError('Empty register')
    return records


def fetch(url):
    if urlparse(url).scheme!='https' or urlparse(url).hostname!='me.gov.ua':raise ValueError('Unexpected source')
    with requests.get(url,timeout=(5,20),stream=True) as r:
        r.raise_for_status()
        if urlparse(r.url).hostname!='me.gov.ua':raise ValueError('Unexpected redirect')
        parts=[];size=0
        for part in r.iter_content(65536):
            size+=len(part)
            if size>60_000_000:raise ValueError('Oversized source')
            parts.append(part)
        return b''.join(parts)


def metadata(kind):
    c=connect()
    try:
        row=c.execute('SELECT metadata FROM snapshots WHERE kind=?',(kind,)).fetchone()
        return json.loads(row[0]) if row else {}
    finally:c.close()


def fresh(meta):
    try:
        age=(datetime.now(timezone.utc)-datetime.fromisoformat(meta['checked_at'])).total_seconds()
        return not meta.get('error') and 0<=age<86400
    except (KeyError,ValueError,TypeError):return False


def refresh(kind,force=False):
    with LOCK:
        old=metadata(kind)
        if not force and fresh(old):return old
        if not force and old.get('last_attempt'):
            if (datetime.now(timezone.utc)-datetime.fromisoformat(old['last_attempt'])).total_seconds()<300:return old
        now=datetime.now(timezone.utc).isoformat()
        try:
            soup=BeautifulSoup(fetch(SOURCES[kind]['index']),'html.parser')
            matches=[urljoin('https://me.gov.ua',a['href']) for a in soup.find_all('a',href=True) if SOURCES[kind]['match'] in normalized(a.get_text(' ',strip=True)) and '/view/' in a['href']]
            if len(set(matches))!=1:raise ValueError('Ambiguous official publication')
            view=matches[0];html=fetch(view).decode('utf-8')
            title=BeautifulSoup(html,'html.parser').title.get_text()
            stamp=re.search(r'(\d{2}\.\d{2}\.\d{4})',title)
            if not stamp:raise ValueError('Missing publication date')
            as_of=datetime.strptime(stamp[1],'%d.%m.%Y').date()
            if as_of>datetime.now(ZoneInfo('Europe/Kyiv')).date():raise ValueError('Future publication date')
            files=re.findall(r'https://me\.gov\.ua/download/[a-zA-Z0-9-]+/file\.xlsx',html)
            if len(set(files))!=1:raise ValueError('Missing workbook')
            url=files[0];content=fetch(url);sha=hashlib.sha256(content).hexdigest()
            rows=parse_book(content,kind)
            archive=config.DATA/'registers'
            archive.mkdir(parents=True,exist_ok=True)
            (archive/(sha+'.xlsx')).write_bytes(content)
            meta={'kind':kind,'title':SOURCES[kind]['title'],'as_of':as_of.isoformat(),'checked_at':now,'last_attempt':now,'source_url':url,'snapshot_url':'/registers/'+sha+'/xlsx','publication_url':view,'sha256':sha,'rows':len(rows),'error':None}
            c=connect()
            try:
                with c:
                    c.execute('DELETE FROM records WHERE kind=?',(kind,))
                    c.executemany('INSERT INTO records VALUES (?,?,?,?,?,?)',rows)
                    c.execute('INSERT OR REPLACE INTO snapshots VALUES (?,?)',(kind,json.dumps(meta,ensure_ascii=False)))
            finally:c.close()
            return meta
        except Exception as exc:
            meta={**old,'kind':kind,'title':SOURCES[kind]['title'],'last_attempt':now,'error':type(exc).__name__}
            c=connect()
            try:
                with c:c.execute('INSERT OR REPLACE INTO snapshots VALUES (?,?)',(kind,json.dumps(meta,ensure_ascii=False)))
            finally:c.close()
            return meta


def row_status(kind,payload):
    if kind=='varieties':return 'Виключений із реєстру' if payload.get('section')=='excluded' else 'Є в основному аркуші оприлюдненого реєстру'
    if payload.get('cancelled') or re.search(r'ануль|скасован|призупин',payload.get('notes',''),re.I):return 'Є позначка про скасування/призупинення — потрібна перевірка'
    try:
        end=date.fromisoformat(payload.get('expires',''))
        return 'Зазначений строк минув' if end<datetime.now(ZoneInfo('Europe/Kyiv')).date() else 'Зазначений строк ще не минув; це не перевірка етикетки'
    except ValueError:return 'Строк потребує ручної перевірки'


def record_priority(kind, payload):
    if kind=='varieties':return 2 if payload.get('section')=='excluded' else 0
    status=row_status(kind,payload)
    if 'скасування' in status or 'строк минув' in status:return 2
    if 'ще не минув' in status:return 0
    return 1


def search(kind,query,update=True):
    if kind not in SOURCES or not 2<=len(query.strip())<=150:raise ValueError('Потрібні тип реєстру та назва від 2 до 150 символів.')
    meta=refresh(kind) if update else metadata(kind)
    q=normalized(query).replace('\\','\\\\').replace('%','\\%').replace('_','\\_')
    c=connect()
    try:
        terms=q.split()
        clause="kind=? AND "+" AND ".join("search LIKE ? ESCAPE '\\'" for _ in terms)
        values=(kind, *('%'+term+'%' for term in terms))
        total=c.execute('SELECT count(*) FROM records WHERE '+clause,values).fetchone()[0]
        rows=c.execute("SELECT * FROM records WHERE "+clause+" ORDER BY CASE WHEN json_extract(payload,'$.section')='excluded' THEN 1 ELSE 0 END,name,row_no DESC",values).fetchall()
    finally:c.close()
    items=[]
    for r in rows:
        p=json.loads(r['payload']);items.append({'sheet':r['sheet'],'row':r['row_no'],**p,'assessment':row_status(kind,p),'historical':record_priority(kind,p)==2})
    items.sort(key=lambda p:(record_priority(kind,p),normalized(p['name']),-p['row']))
    historical_total=sum(p['historical'] for p in items)
    return {'historical_total':historical_total,'shown':min(len(items),20),'metadata':meta,'fresh_check':fresh(meta),'query':query,'total':total,'items':items[:20],'note':'Це пошук записів: він показує назву, склад, номер посвідчення, зазначений строк і позначки скасування. Запис без прострочення не означає, що препарат підходить для вашого поля. Дозу, суміш, дату обробки та справжність каністри ця функція не перевіряє. Відсутність збігу не доводить заборону. Записи з однаковою назвою можуть відрізнятися; звірте номер, склад і форму. Для застосування потрібна чинна українська етикетка та умови поля. Реєстрація сорту не означає найкращу врожайність для вашого поля.'}


def answer(kind,query):
    result=search(kind,query);m=result['metadata']
    lines=[f"{SOURCES[kind]['title']}: версія {m.get('as_of','не встановлена')}. Перевірка джерела: {m.get('checked_at','не виконана')}."]
    if not result['fresh_check']:lines.append('Свіжу перевірку не підтверджено. Нижче — лише збережена копія, якщо вона є.')
    lines.append(f"Знайдено {result['total']} записів, з них історичних/скасованих: {result['historical_total']}. Це можуть бути різні посвідчення одного найменування. Спочатку показано записи без позначки скасування, строк яких ще не минув; невідомі строки потребують перевірки. Деталі — Господарство → Реєстри.")
    primary=[item for item in result['items'] if not item['historical']]
    if not primary and result['items']:
        lines.append('Знайдені лише історичні/скасовані записи. Нижче один приклад; це не підтвердження актуальної реєстрації.')
    for item in (primary[:5] if primary else result['items'][:1]):
        details=[f"{item['name']} — {item['assessment']}",f"Номер: {item.get('number','—')}; аркуш «{item['sheet']}», рядок {item['row']}"]
        for k,label in [('ingredient','Склад'),('crop','Культура'),('expires','Строк у реєстрі'),('zone','Зона (код джерела)'),('notes','Примітка')]:
            if item.get(k):details.append(label+': '+item[k])
        lines.append('\n'.join(details))
    lines.extend([result['note'],'Офіційне джерело: '+m.get('source_url',SOURCES[kind]['index'])])
    return {'answer':'\n\n'.join(lines),'pages':[],'sources':[],'status':'registry_lookup','registry':result}


def query_intent(question):
    m=re.fullmatch(r'\s*(?:перевір|знайди)\s+(препарат|сорт)\s*:?\s+(.+?)\s*',question,re.I)
    if not m:
        m=re.fullmatch(r'\s*чи\s+(?:зареєстрован(?:ий|а)|дозволен(?:ий|а))\s+(препарат|сорт)\s+(.+?)\s*\??',question,re.I)
    if not m or not 2<=len(m[2])<=150:return None
    return ('pesticides' if m[1].lower()=='препарат' else 'varieties',m[2].strip(' «»?'))
